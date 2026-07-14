"""Imports trucker leads from the FMCSA-sourced Excel into the dispatch CRM.
Idempotent: skips DOT numbers already present. Run locally AND once in
Render's Web Shell (the xlsx ships in the repo under data/ for that reason)."""

from pathlib import Path

from openpyxl import load_workbook

from app import create_app
from app.models import db, TruckerLead

XLSX = Path(__file__).parent / "data" / "LEADS_TRUCKERS_TX.xlsx"

app = create_app()
with app.app_context():
    existing = {l.dot_number for l in TruckerLead.query.all() if l.dot_number}
    wb = load_workbook(XLSX)
    added = 0
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for row in ws.iter_rows(min_row=2, values_only=True):
            company, city, phone, mc, dot, alta = row[0], row[1], str(row[2] or ""), str(row[3] or "").replace("MC-", ""), str(row[4] or ""), str(row[5] or "")
            if not company or dot in existing:
                continue
            db.session.add(TruckerLead(
                company_name=company, city=city, state="TX", phone=phone,
                mc_number=mc, dot_number=dot, authority_date=alta,
                notes=f"Fuente: FMCSA census ({sheet})",
            ))
            existing.add(dot)
            added += 1
    db.session.commit()
    print(f"Importados: {added}. Total truckers: {TruckerLead.query.count()}")
